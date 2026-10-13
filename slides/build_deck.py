"""Build the findings deck for the Privacy Externality Game (one finding per slide).

Each slide: what we looked at, what we found, how much to trust it (A/B/C,
see FINDINGS_MAP.md). Attack and sorting numbers are read from the
post-solver-fix caches; figures come from paper/figures and slides/make_figs.py. Run with python-pptx on the path:

    PYTHONPATH=<dir with python-pptx> .venv/bin/python slides/build_deck.py
"""
from __future__ import annotations

import pickle
import re
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_MARKER_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
ASSETS = Path(__file__).resolve().parent / "assets"
CACHE = ROOT / "game_experiments" / "_cache_v2"
OUT = Path(__file__).resolve().parent / "peg_findings.pptx"

NAVY = RGBColor(0x08, 0x30, 0x6B)
ORANGE = RGBColor(0xFD, 0x8C, 0x3B)
GREY = RGBColor(0x4D, 0x4D, 0x4D)
LIGHT = RGBColor(0xF2, 0xF4, 0xF7)
RULE = RGBColor(0xD0, 0xD6, 0xDE)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Calibri"

# ---------------------------------------------------------------- data
c4 = pickle.load(open(CACHE / "C4.pkl", "rb"))["rows"]
h4 = pickle.load(open(CACHE / "H4.pkl", "rb"))["per_k"]
passive = {r["k"]: r["A_target"] for r in c4}
aware = {r["k"]: r["A_target"] for r in h4}
collapse = [passive[k] / aware[k] for k in (1, 2, 3, 4)]
c8 = pickle.load(open(CACHE / "C8_K200.pkl", "rb"))["rows"]
sorting = []
for r in c8:
    if r["ratio"] == 1.0:
        continue
    ok = sum(1 for i, j in r["nes"] if (i >= j if r["ratio"] < 1 else i <= j))
    sorting.append((r["ratio"], len(r["nes"]), ok))
tot, okt = sum(n for _, n, _ in sorting), sum(o for _, _, o in sorting)
print("collapse factors", [round(x, 1) for x in collapse])
print("sorting", sorting)

# ---------------------------------------------------------------- helpers
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]
TOKEN = re.compile(r"(\*\*.+?\*\*|_\{.+?\}|\^\{.+?\})")


def runs(paragraph, text, size, color=GREY, bold=False):
    """Add text with **bold**, _{sub} and ^{sup} markup."""
    for part in TOKEN.split(text):
        if not part:
            continue
        r = paragraph.add_run()
        f = r.font
        f.name, f.size, f.color.rgb, f.bold = FONT, Pt(size), color, bold
        if part.startswith("**"):
            r.text, f.bold = part[2:-2], True
            f.color.rgb = NAVY
        elif part.startswith("_{") or part.startswith("^{"):
            r.text = part[2:-1]
            r.font._rPr.set("baseline", "-25000" if part[0] == "_" else "30000")
        else:
            r.text = part


def textbox(slide, x, y, w, h, items, size=13, color=GREY, bullet=True,
            space=5, anchor=MSO_ANCHOR.TOP, align=PP_ALIGN.LEFT, bold=False):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    for n, item in enumerate(items):
        p = tf.paragraphs[0] if n == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space)
        level = 0
        if isinstance(item, tuple):
            level, item = item
        if bullet:
            mark = "•  " if level == 0 else "–  "
            r = p.add_run()
            r.text = mark
            r.font.name, r.font.size = FONT, Pt(size)
            r.font.color.rgb = ORANGE if level == 0 else GREY
            r.font.bold = True
            if level:
                p.level = 1
        runs(p, item, size - (1 if level else 0), color, bold)
    return tb


def header(slide, kicker, title, n):
    textbox(slide, 0.55, 0.32, 12.2, 0.35, [kicker.upper()], size=12,
            color=ORANGE, bullet=False, bold=True)
    textbox(slide, 0.55, 0.62, 12.3, 0.95, [title], size=26, color=NAVY,
            bullet=False, bold=True)
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(0.6),
                                      Inches(1.55), Inches(12.73), Inches(1.55))
    line.line.color.rgb = RULE
    line.line.width = Pt(1)
    textbox(slide, 0.55, 7.05, 8, 0.3,
            ["The Privacy Externality Game · strategic budgets in sampling-based iDP"],
            size=9, color=RGBColor(0x8A, 0x8A, 0x8A), bullet=False)
    textbox(slide, 12.0, 7.05, 0.8, 0.3, [str(n)], size=9,
            color=RGBColor(0x8A, 0x8A, 0x8A), bullet=False, align=PP_ALIGN.RIGHT)


def card(slide, x, y, w, h, title, items, size=11.5, fill=LIGHT):
    box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y),
                                 Inches(w), Inches(h))
    box.fill.solid()
    box.fill.fore_color.rgb = fill
    box.line.fill.background()
    box.shadow.inherit = False
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y),
                                 Inches(0.06), Inches(h))
    bar.fill.solid()
    bar.fill.fore_color.rgb = NAVY
    bar.line.fill.background()
    bar.shadow.inherit = False
    textbox(slide, x + 0.2, y + 0.12, w - 0.3, 0.3, [title.upper()], size=11,
            color=NAVY, bullet=False, bold=True)
    textbox(slide, x + 0.2, y + 0.45, w - 0.3, h - 0.5, items, size=size, space=4)


def picture(slide, path, x, y, w, h):
    from PIL import Image
    iw, ih = Image.open(path).size
    scale = min(w / iw, h / ih)
    pw, ph = iw * scale, ih * scale
    return slide.shapes.add_picture(str(path), Inches(x + (w - pw) / 2),
                                    Inches(y + (h - ph) / 2), Inches(pw), Inches(ph))


def caption(slide, x, y, w, text):
    textbox(slide, x, y, w, 0.5, [text], size=10, color=RGBColor(0x6A, 0x6A, 0x6A),
            bullet=False)


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def table(slide, x, y, w, h, rows, col_w=None, size=11, header_fill=NAVY):
    shape = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y),
                                   Inches(w), Inches(h))
    tbl = shape.table
    if col_w:
        for i, cw in enumerate(col_w):
            tbl.columns[i].width = Inches(cw)
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = tbl.cell(i, j)
            cell.margin_left = cell.margin_right = Inches(0.06)
            cell.margin_top = cell.margin_bottom = Inches(0.03)
            tf = cell.text_frame
            tf.paragraphs[0].alignment = PP_ALIGN.CENTER if j else PP_ALIGN.LEFT
            runs(tf.paragraphs[0], str(val), size, WHITE if i == 0 else GREY,
                 bold=(i == 0))
            cell.fill.solid()
            cell.fill.fore_color.rgb = header_fill if i == 0 else (
                LIGHT if i % 2 else WHITE)
    return tbl


def style_chart(chart, size=11):
    chart.font.name = FONT
    chart.font.size = Pt(size)
    chart.font.color.rgb = GREY


def log_axis(axis):
    scaling = axis._element.find(
        "{http://schemas.openxmlformats.org/drawingml/2006/chart}scaling")
    lb = etree.SubElement(
        scaling, "{http://schemas.openxmlformats.org/drawingml/2006/chart}logBase")
    lb.set("val", "10")
    scaling.insert(0, lb)


TRUST = {
    "A": ("A · proved / exact", RGBColor(0x2E, 0x7D, 0x32)),
    "B": ("B · solid numerics", RGBColor(0x15, 0x65, 0xC0)),
    "C": ("C · provisional", RGBColor(0xE6, 0x8A, 0x00)),
}


def badge(slide, x, y, level):
    label, col = TRUST[level]
    b = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y),
                               Inches(1.75), Inches(0.3))
    b.adjustments[0] = 0.5
    b.fill.solid()
    b.fill.fore_color.rgb = col
    b.line.fill.background()
    b.shadow.inherit = False
    tf = b.text_frame
    tf.margin_top = tf.margin_bottom = Inches(0)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    runs(p, label, 10.5, WHITE, bold=True)


def block(slide, x, y, w, h, title, items, size=12.5):
    textbox(slide, x, y, w, 0.3, [title], size=10.5, color=ORANGE, bullet=False,
            bold=True)
    textbox(slide, x, y + 0.3, w, h - 0.3, items, size=size, space=5)


PLAIN_FILL = RGBColor(0xFF, 0xF1, 0xE6)


def plain_band(slide, x, y, w, h, text, size=11):
    """'In plain words' interpretation for a non-specialist reader."""
    box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y),
                                 Inches(w), Inches(h))
    box.fill.solid()
    box.fill.fore_color.rgb = PLAIN_FILL
    box.line.fill.background()
    box.shadow.inherit = False
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y),
                                 Inches(0.06), Inches(h))
    bar.fill.solid()
    bar.fill.fore_color.rgb = ORANGE
    bar.line.fill.background()
    bar.shadow.inherit = False
    textbox(slide, x + 0.18, y + 0.06, w - 0.3, 0.28, ["IN PLAIN WORDS"], size=10.5,
            color=ORANGE, bullet=False, bold=True)
    textbox(slide, x + 0.18, y + 0.33, w - 0.3, h - 0.38, [text], size=size,
            color=GREY, bullet=False)


def finding(n, kicker, title, fig, looked, found, trust_levels, trust, files,
            fig_w=7.0, caption_text=None, plain=None):
    s = prs.slides.add_slide(BLANK)
    n = len(prs.slides)
    header(s, f"Finding {n - 1} · {kicker}", title, n)
    fx, fy, fh = 0.55, 1.72, 3.62
    picture(s, ASSETS / fig, fx, fy, fig_w, fh)
    if caption_text:
        caption(s, fx, 5.34, fig_w, caption_text)
    if plain:
        plain_band(s, fx, 5.66, fig_w, 1.32, plain)
    x = fx + fig_w + 0.3
    w = 13.333 - x - 0.5
    block(s, x, 1.72, w, 1.55, "WHAT WE LOOKED AT", looked, size=11.5)
    block(s, x, 3.3, w, 1.95, "WHAT WE FOUND", found, size=12.5)
    textbox(s, x, 5.3, w, 0.3, ["HOW MUCH TO TRUST IT"], size=10.5, color=ORANGE,
            bullet=False, bold=True)
    for k, lv in enumerate(trust_levels):
        badge(s, x + k * 1.85, 5.62, lv)
    textbox(s, x, 5.98, w, 0.95, trust, size=11, space=3)
    notes(s, "Files: " + files)
    return s


# ================================================================ 1 title
s = prs.slides.add_slide(BLANK)
bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
bg.fill.solid()
bg.fill.fore_color.rgb = NAVY
bg.line.fill.background()
stripe = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.75),
                            Inches(1.2), Inches(0.08))
stripe.fill.solid()
stripe.fill.fore_color.rgb = ORANGE
stripe.line.fill.background()
textbox(s, 0.75, 1.95, 11.8, 1.0, ["The Privacy Externality Game"], size=42,
        color=WHITE, bullet=False, bold=True)
textbox(s, 0.75, 2.85, 11.8, 0.7,
        ["Strategic budget choices in sampling-based individual differential privacy"],
        size=22, color=RGBColor(0xC9, 0xD6, 0xEA), bullet=False)
textbox(s, 0.75, 3.5, 11.8, 0.9, [
    "Silos declare ε_{i}; the mechanism solves one σ and rates p_{i} (every budget binds, "
    "Σ|G_{i}|p_{i} = B); each silo maximises  u_{i} = V(σ) − λ_{i} A_{PLD,i}.",
    "Canonical instance: 2 silos × 5000 records, B = 128, T = 390, δ = 10^{−12}, "
    "ε ∈ [0.5, 64], λ = 1.",
], size=13.5, color=RGBColor(0xE6, 0xEC, 0xF5), bullet=False, space=4)
textbox(s, 0.75, 4.45, 11.8, 0.3, ["IN PLAIN WORDS"], size=11, color=ORANGE, bullet=False,
        bold=True)
textbox(s, 0.75, 4.72, 11.8, 0.7, [
    "With individual differential privacy, each data owner picks how private their data should be. "
    "In the popular sampling-based version, all choices are tied together through one shared batch and "
    "one shared noise level, so your privacy depends on what everyone else picks. We ask what happens "
    "when everyone picks rationally."], size=13, color=WHITE, bullet=False)
textbox(s, 0.75, 5.6, 4, 0.3, ["EVERY SLIDE: ONE FINDING, WITH A TRUST LEVEL"], size=11,
        color=ORANGE, bullet=False, bold=True)
for k, (lv, desc) in enumerate([
        ("A", "analytic proof or exact accounting (closed form / certified)"),
        ("B", "post-fix deployed-accountant numerics, K = 200 grid"),
        ("C", "pre-fix, coarse, or single-configuration evidence")]):
    badge(s, 0.8, 5.95 + k * 0.36, lv)
    textbox(s, 2.7, 5.93 + k * 0.36, 9, 0.35, [desc], size=12,
            color=RGBColor(0xE6, 0xEC, 0xF5), bullet=False)

# ================================================================ 2 externality
finding(2, "the externality",
        "Your risk is set by the others' budgets, not by your contract",
        "fig_C5_surface.png",
        ["Target budget fixed (ε_{T} = 4, 16); other group's budget ε_{O} and data "
         "share (0.2–0.8) swept; A_{PLD} by exact PLD accounting."],
        ["Same contract: A_{PLD} falls **0.20 → ≈0** (ε_{T}=4) and **0.55 → 0.03** "
         "(ε_{T}=16) as the others loosen.",
         "Curves cross at ε_{O} ≈ ε_{T}: a large foreign share amplifies the effect "
         "below, damps it above.",
         "Tightening any budget raises everyone else's A_{PLD}."],
        ["A", "B"],
        ["Sign proved for small p. Full K = 200 grid: 0 violations in 25,612 unit "
         "differences (321 ties). No analytic proof beyond small p."],
        "run_continuous.py (C5) → _cache_v2/C5.pkl; tasks_findings/20, 11; FINDINGS_MAP T1",
        fig_w=7.4, caption_text="Paper Fig. 15. Canonical instance, deployed RDP accountant.",
        plain='Each training batch has a fixed number of seats. If the other group asks for more privacy, it takes fewer seats, so your records are drawn more often to fill the batch, and more noise is added. The extra sampling outweighs the extra noise: your data leaks more, although your own setting never changed.')

# ================================================================ 3 frontier
finding(3, "equilibrium structure",
        "Identical players mostly split into a hiding and a carrying role",
        "fig_frontier_ne.png",
        ["Symmetric 2-player game, λ = 1: all pure NE on the K = 200 deployed grid; "
         "exact first-order conditions at the equilibria."],
        ["42 of 43 grid NE sit **on the participation frontier**: tight player at the "
         "floor rate, loose player 6.4–6.7× looser.",
         "Risk **≈7·10^{−4} vs 0.21–0.32**; the hiding seat earns more "
         "(0.43–0.59 vs 0.23–0.27): Hawk–Dove.",
         "Corner equilibria (two one-sided conditions): the grid samples **one "
         "continuum**. The 43rd point is the symmetric equilibrium (next slide)."],
        ["A", "B"],
        ["Corner conditions exact (loose player pushes into the wall, +9.7·10^{−3}); "
         "the asymmetric family survives every grid, V and accountant tested."],
        "run_locus.py (C6) → C6_K200.pkl; exact_foc.py; slides/make_figs.py; FINDINGS_MAP T2",
        fig_w=7.0,
        plain='Two identical groups do not end up behaving alike. One sets a strict budget and is almost never sampled (it hides at the floor); the other fills nearly all batch seats. The hider gets the good model almost for free. As in Hawk–Dove, both prefer the hiding role, but only one can have it.')

# ================================================================ symmetric
finding(0, "symmetric moderation",
        "There is one symmetric equilibrium, and it needs composition",
        "fig_t390_local_slice.png",
        ["Exact composition (no accountant), T = 390, λ = 1: player 2's utility on 14 "
         "dense budgets around the diagonal for ε_{1} = 2.548 and 2.92 (Slurm 292170)."],
        ["The own slice is a **peak**, not a trough: BR(2.548) ≈ 2.62, BR(2.92) ≈ 2.38.",
         "Symmetric NE at **ε* ≈ 2.59** (σ ≈ 1.16 < √3); BR slope ≈ −0.65, "
         "so best-response play is locally stable.",
         "T ≤ 2: no interior symmetric NE at any λ. The deployed RDP grid shows this "
         "point at 2.96, as a 10^{−4} near-tie."],
        ["A", "B"],
        ["Exact profiles, dense near the diagonal. ε* is bracketed (2.55, 2.92) by two "
         "columns; stability from a two-point slope. The earlier 21-point columns had "
         "no deviation within ±40% of the diagonal."],
        "game_experiments/t390_local_slice.py → _scripts/83_t390_local_slice.pkl; multi_step_game.py; FINDINGS_MAP T4",
        fig_w=6.9,
        plain='There is one exception: if both pick a moderate budget of about 2.6, neither gains by moving. It exists only because training repeats the noisy step many times; with 1–2 steps it does not. The peak is very flat, so a coarse search can easily miss it or invent it.')

# ================================================================ V-shape
finding(0, "utility model",
        "Utility shape moves the symmetric point, never the asymmetric family",
        "fig_v_shape_postfix.png",
        ["Six admissible V(σ): 1/(1+σ^{2}), 1/(1+σ), e^{−σ²/2}, e^{−σ}, −log(1+σ), −σ; "
         "pure NE re-enumerated on the same K = 200 mechanism grid (V-free)."],
        ["All asymmetric NE lie on the **same participation frontier**; V only slides "
         "their extent (27–45 NE).",
         "Five of six V show one symmetric point, at a V-dependent budget "
         "(1.4–4.2); linear V = −σ shows none.",
         "So the symmetric point is a property of V's shape; the asymmetric family is a property of the mechanism."],
        ["B", "C"],
        ["Frontier family: solid. Symmetric points: deployed-grid near-ties "
         "(margins ≲ 10^{−4}), exact check done only for V_{1}."],
        "tasks_findings/_scripts/15_v_shape.py, 82_paper_audit.py; slides/make_figs.py; FINDINGS_MAP T15",
        fig_w=7.4, caption_text="Paper Fig. 14, recomputed on the post-fix grid. Hollow: "
                                "symmetric grid points.",
        plain="We don't know exactly how model quality depends on noise, so we tried six plausible curves. The 'one hides, one carries' outcomes barely move: they come from how seats are allocated, not from the curve. Only the moderate equal-budget point depends on it, and a straight-line curve removes it.")

# ================================================================ 4 corners
finding(4, "why",
        "A budget buys participation before it buys model quality",
        "t1_slice_decomposition.png",
        ["Exact T = 1 game (closed form): own-budget slice split into value gain and "
         "privacy cost; λ swept over 10^{−3}–80; group size 500–50,000."],
        ["**Rate phase:** exposure takes off ~e^{ε} while σ is still pinned by the "
         "opponent → a trough. Moderate participation is the worst reply.",
         "**Noise phase:** rate capped; extra budget lowers σ.",
         "Best replies: floor, frontier, or noise phase, hence the corner equilibria."],
        ["A", "B"],
        ["Machine precision at T = 1, certified at T = 2. Across 191 robustness worlds "
         "corner play holds in 107/110 linear-cost ones; convex cost breaks it."],
        "single_step_*.py, two_step_*.py, _scripts/46, 66; nummeric_eval §7–15; FINDINGS_MAP T3",
        fig_w=7.4, caption_text="T = 1, λ = 0.7, |G| = 5000, B = 128; opponent fixed at "
                                "ε_{1} = 29.85.",
        plain='Raising your budget first mainly gets you sampled more, so your exposure jumps while the other group still controls the noise and the model hardly improves. Only once you hold most of the batch does extra budget make the model cleaner. The middle is the worst spot: stay strict or go all in.')

# ================================================================ 6 common weight
finding(6, "shared privacy preference",
        "Caring more about privacy shifts the asymmetric family down",
        "fig_C7_lambda_ratio.png",
        ["Common weight λ ∈ {0.125, …, 4} for both players; all pure NE on the "
         "K = 200 deployed grid."],
        ["λ ≤ 0.5: 34–62 NE, all asymmetric; λ = 1 adds the symmetric point.",
         "Higher λ pulls the whole family to lower budgets.",
         "λ = 2: no pure NE on this grid; λ = 4: single corner (0.5, 0.5)."],
        ["B", "C"],
        ["Asymmetric family: solid. The λ = 2 gap is accountant-conditional: with "
         "761 Rényi orders the cycle collapses onto pure NE."],
        "run_locus.py (C7) → C7_K200.pkl; run_cycle_zoom.py; FINDINGS_MAP T6",
        fig_w=6.2, caption_text="Paper Fig. 16. Crosses: centroids; ellipses: 68% "
                                "dispersion.",
        plain='If both groups care more about privacy, both choose stricter budgets, but the split into one hider and one carrier remains. Turning up how much everyone values privacy does not make the groups behave the same way.')

# ================================================================ 7 sorting
finding(7, "heterogeneity",
        "Preferences decide who is exposed: whoever values privacy least",
        "fig_C8_asymmetric_lambda.png",
        ["λ_{1}/λ_{2} ∈ {0.25, 0.5, 1, 2, 4}, K = 200; group-size ratios 0.1–10; "
         "n = 3–10 players with λ_{i} ∈ geomspace(0.25, 4)."],
        [f"**{okt}/{tot}** heterogeneous NE give the lower budget to the more "
         "sensitive player; the cloud jumps sides at λ_{1} = λ_{2}.",
         "Group size does **not** sort: smaller groups appear in both roles.",
         "n = 3–10: budget order is the exact reverse of λ order in every NE."],
        ["A", "B"],
        ["Proved under an exchange condition; λ sweep post-fix. n-player: 100% "
         "sorting up to n = 10 (coarse 7-action grid)."],
        "run_locus.py (C8), run_continuous.py (C2), run_nplayer.py; tasks_findings/07, 19; FINDINGS_MAP T7, T10",
        fig_w=6.2,
        plain='When the groups value privacy differently, the outcome is predictable: the group that cares least ends up carrying the batch and the exposure. Who gets exposed is set by preferences, not chance. Being the smaller group, on its own, does not protect you.')

# ================================================================ 8 mechanism
finding(8, "mechanism, not paradigm",
        "Individual budgets don't cause it; the shared sampling solve does",
        "fig_scale_ne.png",
        ["Same game, same grid, same V and λ, replayed through sensitivity-based iDP "
         "(per-group clipping under shared σ): A_{i} depends on ε_{i} only."],
        ["Sensitivity-based: **exactly one** pure NE, symmetric at (5.73, 5.73).",
         "Same holds for λ ∈ {0.25, 0.5, 1}.",
         "Sampling-based: 43 NE, 42 asymmetric."],
        ["B"],
        ["Post-fix K = 200 caches for both mechanisms. It acts as a control: remove "
         "the coupling and the Hawk–Dove structure disappears."],
        "game/mechanism_scale.py, run_scale.py → _cache_scale/; tasks_findings/18; FINDINGS_MAP T11",
        fig_w=6.0,
        plain="The other common way to give personal budgets changes each group's gradient clipping instead of its sampling rate. There your risk depends only on your own choice, and both groups simply pick the same budget. So the trouble comes from how sampling-based iDP fills the batch, not from personal budgets.")

# ================================================================ 9 welfare
finding(9, "welfare",
        "Selfish play assigns the roles efficiently, but never pays the carrier",
        "fig_welfare.png",
        ["Welfare u_{1}+u_{2} on the K = 200 cache: planner optimum, best pure NE, "
         "enforced equal budgets; r = λ_{2}/λ_{1} ∈ {1, 2, 4, 8}."],
        ["The planner optimum is **itself asymmetric** (3.35, 24.13): composed "
         "exposure is concave in the rate.",
         "Best NE sits ~9% below the planner at every r.",
         "Enforced equal budgets: 6% below the planner at r = 1, **below zero** at r = 8."],
        ["B"],
        ["Cache-only computations on the post-fix grid; NE welfare is grid-resolved. "
         "No equilibrium transfers utility to the carrier, so fairness needs a transfer."],
        "_scripts/59_alt_plays.py → 59_alt_plays_summary.txt; nummeric_eval §22; FINDINGS_MAP T9",
        fig_w=6.6,
        plain='Even a fair-minded planner would let one group carry the batch, because concentrating exposure costs less overall. Selfish play gets fairly close (about 9% worse), but nobody pays the carrier back. Forcing equal budgets on groups that care differently is the worst option.')

# ================================================================ 10 populations
finding(10, "many players",
        "In a 20-silo consortium, everyone free-rides at the strictest budget",
        "fig_e1_population.png",
        ["100 sampled FL consortia: n = 20, |G_{i}| ~ U{2000…12000}, "
         "λ_{i} ~ LogNormal(0, 0.7), B = 256, actions {10, …, 50}; sequential BR."],
        ["**100/100** converge to the uniform floor (10, …, 10).",
         "One silo moves σ by ~1/20, so the free-riding threshold is λ* ≈ 0.02–0.04.",
         "Welfare **31%** below a sorted planner (planner Gini 0.58)."],
        ["C"],
        ["Pre-solver-fix run, approximate planner, 6-point grid. The run at the "
         "critical weight scale was never analysed."],
        "game_experiments/run_flpop_e1.py, flpop_common.py → _cache_flpop/E1_pop*.json; tasks_findings/28; FINDINGS_MAP T10",
        fig_w=5.6,
        plain='With 20 groups, each one barely affects the shared noise, so no group has a reason to loosen its budget for the common good. Everyone picks the strictest option and hopes others carry the batch: classic free-riding, about 30% worse than a coordinated plan. Preliminary result.')

# ================================================================ 11 attacks
finding(11, "attacks",
        "The attack is coalition play, and fails against an aware target",
        "fig_h4_aware_target.png",
        ["n = 5, others at ε = 32; a k-coalition commits to one budget maximising "
         "A_{PLD,t} (κ → ∞); target passive or best-responding (14-point grids)."],
        [f"Passive: A_{{PLD,t}} {passive[0]:.2f} → **{passive[4]:.2f}** at k = 4; the "
         "coalition plays the participation-preserving minimum, as predicted.",
         f"Aware: **{aware[2]:.3f}–{aware[4]:.3f}**, "
         f"{min(collapse):.0f}–{max(collapse):.0f}× lower, below the no-coalition "
         "baseline."],
        ["A", "B"],
        ["Corollary proved. Post-fix caches. Coarse 14-point action grids, so "
         "ratios are approximate."],
        "run_continuous.py (C4), run_h4.py → C4.pkl, H4.pkl; tasks_findings/13; FINDINGS_MAP T12",
        fig_w=6.2, caption_text="Hollow marker: interior optimum at k = 4 (filled: "
                                "poisoning corner, ε_{C} = 0.5).",
        plain='In the known attack, some groups pick strict budgets on purpose to push a target into more exposure. In our model that is simply what colluding players would do. But it only works on a target that stays put; a target that reacts by tightening its own budget escapes almost entirely.')

# ================================================================ 12 measured
finding(12, "measured adversary",
        "At the attack's noise floor, the strategic pressure vanishes",
        "fig_H1_n2_measured_NE.png",
        ["Adult, MLP, 5000 records/silo, 5×5 grid ε ∈ {2, …, 32}; 128 DP-SGD shadow "
         "models per profile; payoffs from measured LiRA + accuracy."],
        ["A_{LiRA} ∈ [0.107, 0.115] (AUC ≤ 0.506) vs A_{PLD} ∈ [0.004, 0.870].",
         "Externality term flat → measured NE at the accuracy corner (16,16)/(32,32); "
         "analytic NE at (2,2)."],
        ["C"],
        ["Pre-fix mechanism (audit: σ ≤ 2.6%, rates ≤ 18.5% off); one dataset. "
         "The above-floor instance is the main missing experiment."],
        "experiment_scripts/mia/04_mia_pergroup.py, _scripts/24_*; _cache_H1; audit_h1_sigma.py; FINDINGS_MAP T13",
        fig_w=7.6,
        plain="Our risk numbers assume the strongest possible attacker. With real trained models and a real attack (LiRA), the attack could barely tell members from non-members here. Then everyone's risk looks the same, the game disappears, and groups simply maximise accuracy.")

# ================================================================ 13 mitigation
finding(13, "mitigation",
        "Fix the contract, not the players: cap the slack",
        "fig_slack_panel.png",
        ["Slack decomposition A_{PLD,i} = A_{εᵢ,δ} − slack_{i}; budget declaration "
         "as a direct mechanism; linear vs convex exposure pricing."],
        ["**Δ-iDP:** slack_{i} ≤ Δ ⇒ no coalition moves A_{PLD,i} by more than Δ "
         "at fixed ε_{i}.",
         "Budget reports are **not DSIC**: the best report depends on the others.",
         "Linear prices only rescale the game; **convex** pricing gives symmetric "
         "interior NE already at T = 1."],
        ["A", "C"],
        ["Δ bound and non-DSIC proved. Convex-pricing result at sample resolution. "
         "Whether the Δ-restricted game is unique/symmetric, and its cost, is open."],
        "tasks_findings/04, 08; nummeric_eval §16; paper/tex/07_mitigation.tex; FINDINGS_MAP T14",
        fig_w=5.6, caption_text="The shaded slack is the movable part of the risk that "
                                "the Δ-clause caps.",
        plain="The problem is that others can move your risk without breaking your contract. The fix is to cap how far they can move it: promise a maximum excess Δ on top of ε. Asking for honest budget reports cannot fix it, and a flat price per unit of risk doesn't help; only a price that grows with the risk does.")

# ================================================================ audit
s = prs.slides.add_slide(BLANK)
header(s, "Audit · paper vs. data", "Where the paper disagrees with the data", len(prs.slides))
table(s, 0.6, 1.75, 12.15, 3.9, [
    ["Paper claim", "Current data", "Verdict"],
    ["One symmetric NE (text) vs. none (Figs. 3, 14)", "one symmetric NE exists: exact ε* ≈ 2.59; grid shows it at 2.96", "captions wrong"],
    ["Symmetric NE confirmed by 25-point refinement at 2.96", "deployed grid: 10^{−4} near-tie, dissolves at 16×", "scoped"],
    ["Fig. 14 counts 44–60 NE", "27–45 NE post-fix", "stale"],
    ["Minority group plays the lower budget", "both orientations for |G_{1}| < |G_{2}|", "falsified"],
    ["Symmetric NE is the welfare-best", "least ΣA, but W = 0.64 vs 0.86 (frontier)", "falsified"],
    ["Exposed player compensated / loose player indifferent", "u = 0.25 vs 0.50; strictly wall-blocked", "falsified"],
    ["Ordinal-potential 4-cycle (K = 22)", "gone; 6-cycle on K = 200 replaces it", "falsified → fixable"],
    ["n-player: 83% sorting at n = 10, A ratio 480 → 75", "100% at all n; ratios 44–330", "falsified"],
    ["Sensitivity appendix, non-DSIC example (2.92 / 1.21)", "7 → 43 NE; 1.82 / 3.43 (direction flips)", "falsified"],
    ["λ = 2: no pure NE", "761 Rényi orders → pure-NE cluster", "scoped"],
    ["BR-PEG regret, H1 measured instance", "pre-fix runs, not re-run", "unverified"],
], col_w=[4.6, 5.45, 2.1], size=11)
plain_band(s, 0.6, 5.85, 12.15, 1.12,
           "The paper grew over many drafts, and a solver bug fixed in July shifted several numbers. "
           "The core story survives: others set your risk, equal groups split into a hider and a "
           "carrier, and the least privacy-minded group is exposed. These rows must be fixed. "
           "Full list with line references: PAPER_AUDIT.md.", size=12)
notes(s, "Holds: externality sign, 43-NE frontier structure, ratio band, Hawk–Dove numbers, "
         "λ-ratio sorting 170/174, PoEV sweep, sensitivity-based control, H4 values.")

prs.save(OUT)
print("saved", OUT)
